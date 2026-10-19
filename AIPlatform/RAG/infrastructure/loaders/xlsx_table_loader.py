"""Læser OT's dokumenter: alle .xlsx-filer i en mappe, hvor hvert ark er en tabel.

Hvorfor en loader ved siden af XlsxLoader?
Testcase-arkene er formularer med user stories og testcases. OT's ark er almindelige tabeller,
fx en signalliste med én række pr. signal. Her bliver hver række sin egen blok
("overskrift: værdi" pr. celle), så ParagraphChunker giver ét afsnit pr. række.
"""
from pathlib import Path

from openpyxl import load_workbook

from RAG.domain.models import Document
from RAG.infrastructure.loaders.xlsx_loader import clean

# Hvorfor stoppe efter mange tomme rækker?
# Et ark kan være formateret langt ned (signallisten siger over en million rækker, men har ca. 12.000).
# At læse alle de tomme rækker ville gøre opstarten langsom.
MAX_EMPTY_ROWS = 1000

# Celler, der ikke siger noget, fx når et opslag i Excel ikke fandt noget.
EMPTY_VALUES = {"", "#N/A"}


class XlsxTableLoader:

    def __init__(self, folder: str | Path):
        self.folder = Path(folder)

    def load(self) -> list[Document]:
        documents = []
        # Også undermapper, ligesom PdfLoader.
        for file in sorted(f for f in self.folder.rglob("*") if f.is_file() and f.suffix.lower() == ".xlsx" and not f.name.startswith("~$")):
            workbook = load_workbook(file, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                text = self._table(sheet.iter_rows(values_only=True))
                if text:
                    documents.append(Document(source=f"{file.relative_to(self.folder).as_posix()} › {sheet.title.strip()}", text=text))
        return documents

    @staticmethod
    def _table(rows) -> str:
        header = None
        blocks = []
        empty_rows = 0
        for row in rows:
            cells = [clean(c) for c in row]
            if not any(v not in EMPTY_VALUES for v in cells):
                empty_rows += 1
                if empty_rows >= MAX_EMPTY_ROWS:
                    break
                continue
            empty_rows = 0

            # Den første række med indhold er overskrifterne.
            if header is None:
                header = cells
                continue

            # Kolonner uden overskrift er hjælpekolonner i regnearket og springes over.
            lines = [f"{h}: {v}" for h, v in zip(header, cells) if h and v not in EMPTY_VALUES]
            if lines:
                blocks.append("\n".join(lines))
        return "\n\n".join(blocks)
