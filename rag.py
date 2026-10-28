"""RAG: indlæser dokumenter (testcases, kravspec) og søger i dem.

Hvorfor egne modeller?
Resten af koden skal kunne tale om "et dokument", "et afsnit" og "et søgeresultat"
uden at vide, om dokumenterne kommer fra filer, SharePoint eller en database,
eller om søgningen sker på ord eller på betydning.
"""
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook


# Hvorfor frozen=True?
# Et dokument eller et afsnit ændrer sig ikke, mens vi arbejder med det.
# Uforanderlige objekter kan trygt deles mellem indlæsning, indeks og svar.
@dataclass(frozen=True)
class Document:
    # Hvorfor gemme kilden?
    # Et svar er kun troværdigt, hvis man kan se, hvor det kommer fra.
    # Kilden følger med hele vejen fra filen til svaret.
    source: str
    text: str


@dataclass(frozen=True)
class Chunk:
    """Et afsnit af et dokument. Det er afsnit, vi søger i, ikke hele dokumenter."""

    # Hvorfor dele dokumenter op?
    # Et helt dokument er for stort at sende med hvert spørgsmål og indeholder
    # meget, der ikke er relevant. Små afsnit giver præcise opslag og billigere kald.
    source: str
    text: str


@dataclass(frozen=True)
class SearchResult:
    chunk: Chunk

    # Hvorfor en score?
    # Søgningen skal kunne sige, hvor godt et afsnit passer, så de bedste
    # kommer først, og helt irrelevante kan sorteres fra.
    score: float


# Hvorfor en port til indlæsning?
# I dag læser vi .md-filer fra en mappe. I morgen kan det være PDF'er,
# SharePoint eller en database. Resten af laget skal ikke mærke forskel.
class DocumentLoader(Protocol):
    def load(self) -> list[Document]:
        """Henter alle dokumenter."""
        ...


# Hvorfor en port til opdeling?
# Hvordan man bedst deler tekst op, afhænger af teksten: afsnit, overskrifter,
# et fast antal ord. Det er noget, man vil eksperimentere med uden at røre resten.
class Chunker(Protocol):
    def split(self, document: Document) -> list[Chunk]:
        """Deler ét dokument op i afsnit."""
        ...


# Hvorfor én port til både at tilføje og søge?
# Det er samme "kartotek": det, der lægges ind, er det, der søges i.
# Skiftes ordsøgning ud med søgning på betydning (embeddings), skiftes begge dele.
class SearchIndex(Protocol):
    def add(self, chunks: list[Chunk]) -> None:
        """Lægger afsnit i indekset."""
        ...

    def search(self, query: str, k: int) -> list[SearchResult]:
        """Finder de k afsnit, der passer bedst til forespørgslen."""
        ...


class IngestService:

    def __init__(self, loader: DocumentLoader, chunker: Chunker, index: SearchIndex):
        # Vi får alle tre dele udefra. Servicen er ligeglad med, hvilke det er.
        self.loader = loader
        self.chunker = chunker
        self.index = index

    def run(self) -> int:
        """Indlæs → del op → læg i indeks. Returnerer antal afsnit."""
        documents = self.loader.load()

        chunks = []
        for document in documents:
            chunks.extend(self.chunker.split(document))

        self.index.add(chunks)
        return len(chunks)


class ParagraphChunker:

    def split(self, document: Document) -> list[Chunk]:
        # Hver user story og hver testcase bliver ét afsnit (en chunk).
        chunks = []
        for paragraph in document.text.split("\n\n"):
            text = paragraph.strip()

            # Spring tomme afsnit og rene overskrifter over. De indeholder ikke svar.
            if not text or text.startswith("#"):
                continue

            chunks.append(Chunk(source=document.source, text=text))
        return chunks


# Små ord, der står i næsten alle sætninger og derfor ikke siger noget om emnet.
STOP_WORDS = {
    "og", "i", "at", "er", "en", "et", "den", "det", "de", "til", "på", "med",
    "for", "som", "har", "man", "hos", "os", "vi", "jeg", "du", "hvor", "hvem",
    "hvad", "hvordan", "skal", "kan", "om", "af", "mange", "der", "ikke",
    # Engelske småord, fordi spørgsmålene stilles på engelsk.
    "the", "a", "an", "is", "are", "does", "do", "our", "we", "my", "it", "of", "to", "in", "and", "or", "what", "how",
}


def words(text: str) -> set[str]:
    """Gør teksten til et sæt af små bogstaver uden stopord."""
    return {w for w in re.findall(r"\w+", text.lower()) if w not in STOP_WORDS}


class KeywordIndex:
    """Simpel søgning: et afsnit scorer ét point for hvert ord, det deler med spørgsmålet.

    Ingen internet, ingen API-nøgle, ingen ekstra pakker.
    Senere kan et embedding-indeks ligge ved siden af og søge på betydning i stedet.
    """

    def __init__(self):
        self.chunks: list[Chunk] = []

    def add(self, chunks: list[Chunk]) -> None:
        self.chunks.extend(chunks)

    def search(self, query: str, k: int) -> list[SearchResult]:
        query_words = words(query)

        results = []
        for chunk in self.chunks:
            score = len(query_words & words(chunk.text))
            if score > 0:
                results.append(SearchResult(chunk=chunk, score=score))

        # Bedste resultater først, og kun de k bedste.
        results.sort(key=lambda r: r.score, reverse=True)
        return results[:k]


class MarkdownLoader:
    """Læser kravspecifikationer fra alle .md-filer i en mappe.

    Hver fil bliver ét dokument. Hvert krav står som sin egen blok adskilt af en tom linje,
    så ParagraphChunker deler dem op med ét krav pr. afsnit, ligesom user stories i xlsx.
    """

    def __init__(self, folder: str | Path):
        self.folder = Path(folder)

    def load(self) -> list[Document]:
        # Hver .md-fil bliver til ét dokument.
        return [
            Document(source=file.name, text=file.read_text(encoding="utf-8"))
            for file in sorted(self.folder.glob("*.md"))
        ]


# En række, der starter en ny blok: US-001, TC-011, US-00x, TC-00x osv.
BLOCK_START = re.compile(r"^(US|TC)-", re.IGNORECASE)

# Området står i filnavnet, fx test-protocol-regninger-mitaalborgforsyning.xlsx.
AREA = re.compile(r"test-protocol-([^-.]+)", re.IGNORECASE)


def area(file: Path) -> str:
    """Området, en fil handler om, fx Regninger. Uden et kendt mønster bruges filnavnet."""
    match = AREA.search(file.name)
    return match.group(1).capitalize() if match else file.name.split(".")[0]


def clean(value) -> str:
    """Gør en celle til én linje tekst. Linjeskift i cellen ville ellers dele blokken op."""
    if value is None:
        return ""
    return " ".join(str(value).split())


class XlsxLoader:
    """Læser user stories og testcases fra alle .xlsx-filer i en mappe.

    Hvert ark bliver ét dokument. Inde i dokumentet står hver user story og hver
    testcase som sin egen blok adskilt af en tom linje, så ParagraphChunker kan
    dele dem op uden at kende til regneark.
    """

    def __init__(self, folder: str | Path):
        self.folder = Path(folder)

    def load(self) -> list[Document]:
        # Hvert ark i hver xlsx-fil bliver til ét dokument med tekst.
        documents = []
        for file in sorted(self.folder.glob("*.xlsx")):
            workbook = load_workbook(file, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = [[clean(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                rows = [row for row in rows if any(row)]
                if not rows:
                    continue

                # Hvorfor to slags ark?
                # "User Stories" er en almindelig tabel med overskrifter.
                # De andre ark er formularer: etiket i kolonne A, værdi i B, testresultater i C og frem.
                if rows[0][0].upper() == "US-ID":
                    text = self._table(rows)
                else:
                    text = self._form(rows)

                # Hvorfor området på hver blok?
                # Hver fil nummererer sine user stories og testcases fra 1, så US-001 findes både
                # under Regninger og Profil. Med området på kan agenten se, hvilken US-001 det er.
                name = area(file)
                text = "\n\n".join(f"Område: {name}\n{block}" for block in text.split("\n\n"))
                source = f"{file.name} › {sheet.title.strip()}"
                documents.append(Document(source=source, text=text))
        return documents

    @staticmethod
    def _table(rows: list[list[str]]) -> str:
        header, *records = rows
        blocks = []
        for record in records:
            lines = [f"{h}: {v}" for h, v in zip(header, record) if h and v]
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks)

    @staticmethod
    def _form(rows: list[list[str]]) -> str:
        blocks: list[list[str]] = []
        for row in rows:
            label, value, *results = row + ["", ""]
            results = [r for r in results if r]

            if BLOCK_START.match(label):
                blocks.append([])
            elif not blocks:
                # Rækker før første blok er testdatoer og bemærkninger øverst i arket.
                if value:
                    blocks.append([value])
                continue

            line = f"{label}: {value}" if label else f"  {value}"
            if results:
                line += f" [test results: {', '.join(results)}]"
            blocks[-1].append(line)
        return "\n\n".join("\n".join(block) for block in blocks)
