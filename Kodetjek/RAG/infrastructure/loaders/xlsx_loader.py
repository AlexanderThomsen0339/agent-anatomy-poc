"""Læser user stories og testcases fra alle .xlsx-filer i en mappe.

Hvert ark bliver ét dokument. Inde i dokumentet står hver user story og hver
testcase som sin egen blok adskilt af en tom linje, så ParagraphChunker kan
dele dem op uden at kende til regneark.
"""
import re
from pathlib import Path

from openpyxl import load_workbook

from RAG.domain.models import Document

# En række, der starter en ny blok: US-001, TC-011, US-00x, TC-00x osv.
BLOCK_START = re.compile(r"^(US|TC)-", re.IGNORECASE)


def clean(value) -> str:
    """Gør en celle til én linje tekst. Linjeskift i cellen ville ellers dele blokken op."""
    if value is None:
        return ""
    return " ".join(str(value).split())


class XlsxLoader:

    def __init__(self, folder: str | Path):
        self.folder = Path(folder)

    def load(self) -> list[Document]:
        # 6: LÆS XLSX: Hvert ark i hver xlsx-fil bliver til ét dokument med tekst.
        documents = []
        for file in sorted(self.folder.glob("*.xlsx")):
            workbook = load_workbook(file, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = [[clean(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                rows = [row for row in rows if any(row)]
                if not rows:
                    continue

                # Hvorfor to slags ark?
                # "User Stories" er en almindelig tabel med overskrifter.
                # De andre ark er formularer: etiket i kolonne A, værdi i B, testresultater i C og frem.
                if rows[0][0].upper() == "US-ID":
                    text = self._table(rows)
                else:
                    text = self._form(rows)

                source = f"{file.name} › {sheet.title.strip()}"
                documents.append(Document(source=source, text=text))
        return documents

    @staticmethod
    def _table(rows: list[list[str]]) -> str:
        header, *records = rows
        blocks = []
        for record in records:
            lines = [f"{h}: {v}" for h, v in zip(header, record) if h and v]
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks)

    @staticmethod
    def _form(rows: list[list[str]]) -> str:
        blocks: list[list[str]] = []
        for row in rows:
            label, value, *results = row + ["", ""]
            results = [r for r in results if r]

            if BLOCK_START.match(label):
                blocks.append([])
            elif not blocks:
                # Rækker før første blok er testdatoer og bemærkninger øverst i arket.
                if value:
                    blocks.append([value])
                continue

            line = f"{label}: {value}" if label else f"  {value}"
            if results:
                line += f" [test results: {', '.join(results)}]"
            blocks[-1].append(line)
        return "\n\n".join("\n".join(block) for block in blocks)
